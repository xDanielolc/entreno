# -*- coding: utf-8 -*-
"""
Convierte los Excel de progresión Bilbo al formato de datos de la app
(esquema versión 1, ver docs/esquema-datos.md).

SOLO LEE los Excel. No escribe, mueve ni modifica ninguno de ellos.

Uso:
    herramientas\\.venv\\Scripts\\python.exe herramientas\\importar_excel.py

Resultado (fuera de Git, porque son datos personales):
    datos-privados\\entrenamiento-importado.json
    datos-privados\\informe-importacion.txt
"""

import datetime as dt
import json
import re
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")  # openpyxl avisa de dibujos que no lee

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------

VERSION_ESQUEMA = 1
DIAS_POR_CICLO = 17
FILA_CABECERA = 2
FILA_PRIMER_DIA = 3

# Carpetas de origen. Son rutas de tu ordenador, así que no se guardan en el
# repositorio: se leen de datos-privados/origenes.json, con esta forma:
#   [{"ruta": "C:\ruta\a\Entrenamientos", "sede": "sede_habitual", "vigente": true}, …]
# «vigente» indica si sus ejercicios son los que usas ahora; una copia antigua
# va con false y solo aporta historial.
ARCHIVO_ORIGENES = Path(__file__).resolve().parent.parent / "datos-privados" / "origenes.json"


def cargar_origenes():
    if not ARCHIVO_ORIGENES.exists():
        sys.exit(f"Falta {ARCHIVO_ORIGENES}. Créalo con las carpetas de tus Excel (ver comentario arriba).")
    origenes = json.loads(ARCHIVO_ORIGENES.read_text(encoding="utf-8"))
    for o in origenes:
        o["ruta"] = Path(o["ruta"])
    return origenes


# Subcarpetas que, estén donde estén, pertenecen a una sede concreta.
SUBCARPETAS_DE_SEDE = {"asturias": "sede_asturias"}

SEDES = [
    {"id": "sede_habitual", "nombre": "Habitual", "archivado": False},
    {"id": "sede_asturias", "nombre": "Asturias", "archivado": False},
]

# Carpetas cuyo contenido ya no se usa.
CARPETAS_ANTIGUAS = ("antiguos", "rutinas y bilbo viejas", "infrecuentes",
                     "completos o ahora no se usan")

# Archivos que no son ejercicios (plantillas y calculadoras).
ARCHIVOS_IGNORADOS = ("herramienta de calculo", "progresion bilbo")

# Carpetas de ejercicios que no dependen de la máquina de un gimnasio.
CARPETAS_UNIVERSALES = ("ejercicios peso corporal",)

GRUPOS = {
    "empuje": "empuje", "push": "empuje",
    "pierna": "pierna", "leg": "pierna", "legs": "pierna",
    "tiron": "tiron", "pull": "tiron",
    "ejercicios peso corporal": "peso-corporal",
}

# Rutina PLPL BILBO-HEAVY DUTY 04-2026, tal como está en el documento de Word.
# Se busca cada fragmento en el nombre del ejercicio (sin tildes ni mayúsculas).
PLPL_SERIES = [
    ("press banca",   [20, 30], "drop-set",           [6, 8]),
    ("fondos",        [20, 30], "drop-set",           [6, 8]),
    ("press militar", [20, 30], "drop-set",           [8, 10]),
    ("crunch",        [20, 30], "excentricas-lentas", [12, 15]),
    ("maquina de abs",[20, 30], "excentricas-lentas", [12, 15]),
    ("sentadilla",    [20, 30], "drop-set",           [6, 8]),
    ("gemelo",        [25, 30], "isometrico-final",   [10, 12]),
    ("abduc",         [20, 30], "excentricas-lentas", [8, 10]),
    ("aduc",          [20, 30], "drop-set",           [8, 10]),
    ("adduc",         [20, 30], "drop-set",           [8, 10]),
    ("remo barra",    [15, 20], "drop-set",           [6, 8]),
    ("remo en barra", [15, 20], "drop-set",           [6, 8]),
    ("pulldown",      [20, 30], "drop-set",           [6, 8]),
    ("dominadas",     [20, 30], "drop-set",           [6, 8]),
    ("upper back",    [20, 30], "drop-set",           [6, 8]),
    ("martillo",      [20, 30], "isometrico-final",   [8, 12]),
    ("pajaros",       [20, 30], "isometrico-final",   [10, 12]),
    ("peso muerto",   [15, 20], "drop-set",           [6, 8]),
    ("isquios",       [20, 30], "drop-set",           [8, 10]),
    ("prensa",        [20, 30], "drop-set",           [8, 10]),
]

# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------


def sin_tildes(texto):
    forma = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in forma if not unicodedata.combining(c))


def normalizar(texto):
    """Minúsculas, sin tildes y con los espacios colapsados."""
    return re.sub(r"\s+", " ", sin_tildes(str(texto)).lower()).strip()


def slug(texto):
    return re.sub(r"[^a-z0-9]+", "_", normalizar(texto)).strip("_")


def nombre_de_archivo(ruta):
    """«1. Press banca 5 (nota).xlsx» -> ("Press banca", "nota")."""
    nombre = ruta.stem
    nombre = re.sub(r"^\d+(\.\d+)?\.?\s*", "", nombre)      # «1. », «2.1 »
    notas = re.findall(r"\(([^)]*)\)", nombre)
    nombre = re.sub(r"\([^)]*\)", "", nombre)
    nombre = re.sub(r"\s+\d+\s*$", "", nombre.strip())        # versión: « 5»
    nombre = nombre.rstrip(". ").strip()
    return nombre[:1].upper() + nombre[1:], "; ".join(n.strip() for n in notas)


def tipo_columna(cabecera):
    """Identifica una columna por su cabecera, ignorando los números que
    Excel añade al convertir un rango en tabla («PESO 4», «Fecha2»)."""
    c = re.sub(r"[\d.]+$", "", normalizar(cabecera or "")).strip(" .")
    if c == "fecha":
        return "fecha"
    if c in ("kg maq", "kg. maq", "kg maq."):
        return "lectura"
    if c == "dia":
        return "dia"
    if c == "peso":
        return "carga"
    if c.startswith("altura"):
        return "carga_altura"
    if c == "repeticiones":
        return "esfuerzo_reps"
    if c == "segundos":
        return "esfuerzo_segundos"
    return None


# ---------------------------------------------------------------------------
# Lectura de un archivo
# ---------------------------------------------------------------------------


def a_fecha(valor):
    """Devuelve (fecha ISO o None, es_aproximada)."""
    if isinstance(valor, (dt.datetime, dt.date)):
        fecha = valor.date() if isinstance(valor, dt.datetime) else valor
        # Una celda con 0 o 1 se lee como 1900: no es una fecha real.
        return (fecha.isoformat(), False) if fecha.year >= 2000 else (None, False)
    if isinstance(valor, str):
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", valor)
        if m:
            d, mes, a = (int(x) for x in m.groups())
            a += 2000 if a < 100 else 0
            try:
                return dt.date(a, mes, d).isoformat(), "aprox" in valor.lower()
            except ValueError:
                pass
    return None, False


def numero(valor):
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return round(float(valor), 3)
    return None


def leer_bloques(hoja):
    """Localiza los bloques de ciclo por las columnas «Fecha» de la cabecera."""
    columnas = [(c, tipo_columna(hoja.cell(FILA_CABECERA, c).value))
                for c in range(1, hoja.max_column + 1)]
    inicios = [c for c, t in columnas if t == "fecha"]
    bloques = []
    for i, inicio in enumerate(inicios):
        fin = inicios[i + 1] if i + 1 < len(inicios) else hoja.max_column + 1
        mapa = {}
        for c, t in columnas:
            if inicio <= c < fin and t and t not in mapa:
                mapa[t] = c
        if "carga" in mapa or "carga_altura" in mapa:
            bloques.append(mapa)
    return bloques


def leer_archivo(ruta):
    """Devuelve la información de un Excel de ciclo, o None si no lo es."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    if "Ciclo" not in libro.sheetnames:
        return None
    hoja = libro["Ciclo"]

    ciclos = []
    for mapa in leer_bloques(hoja):
        col_carga = mapa.get("carga") or mapa.get("carga_altura")
        col_esf = mapa.get("esfuerzo_reps") or mapa.get("esfuerzo_segundos")
        escalera, registros = [], []
        for fila in range(FILA_PRIMER_DIA, FILA_PRIMER_DIA + DIAS_POR_CICLO):
            dia = fila - FILA_PRIMER_DIA + 1
            carga = numero(hoja.cell(fila, col_carga).value)
            escalera.append(carga)
            esfuerzo = numero(hoja.cell(fila, col_esf).value) if col_esf else None
            if not esfuerzo:
                continue
            fecha, aprox = a_fecha(hoja.cell(fila, mapa["fecha"]).value)
            lectura = numero(hoja.cell(fila, mapa["lectura"]).value) if "lectura" in mapa else None
            registros.append({"dia": dia, "fecha": fecha, "fechaAproximada": aprox,
                              "carga": carga, "lectura": lectura, "esfuerzo": esfuerzo})
        ciclos.append({
            "esfuerzo": "tiempo" if "esfuerzo_segundos" in mapa else "repeticiones",
            "carga": "altura" if "carga_altura" in mapa else "peso",
            "lectura": "lectura" in mapa,
            "escalera": escalera,
            "registros": registros,
        })

    # Notas sueltas debajo de la tabla (p. ej. «100kg X 2 20/5/25 hay grabación»).
    notas = []
    for fila in range(FILA_PRIMER_DIA + DIAS_POR_CICLO + 4, hoja.max_row + 1):
        for c in range(1, hoja.max_column + 1):
            v = hoja.cell(fila, c).value
            if isinstance(v, str) and v.strip():
                notas.append(v.strip())

    libro.close()
    return {"ciclos": ciclos, "notas": notas}


# ---------------------------------------------------------------------------
# Recorrido de carpetas
# ---------------------------------------------------------------------------


def sede_y_contexto(origen, ruta):
    partes = [normalizar(p) for p in ruta.relative_to(origen["ruta"]).parts[:-1]]
    sede = origen["sede"]
    if partes and partes[0] in SUBCARPETAS_DE_SEDE:     # p. ej. dentro de una copia antigua
        sede = SUBCARPETAS_DE_SEDE[partes[0]]
    universal = any(p in CARPETAS_UNIVERSALES for p in partes)
    antiguo = any(p in CARPETAS_ANTIGUAS for p in partes)
    grupo = None
    for p in partes:
        clave = re.sub(r"^\d+\.\s*", "", p)
        clave = re.sub(r"\s+\d+$", "", clave)
        if clave in GRUPOS:
            grupo = GRUPOS[clave]
    return {
        "sede": None if universal else sede,
        "vigente": origen["vigente"] and not antiguo,
        "grupo": grupo,
        "carpeta": str(ruta.parent.relative_to(origen["ruta"])),
    }


def recorrer():
    for origen in cargar_origenes():
        if not origen["ruta"].exists():
            print(f"AVISO: no existe {origen['ruta']}")
            continue
        for ruta in sorted(origen["ruta"].rglob("*.xlsx")):
            if ruta.name.startswith("~$"):
                continue
            if any(ign in normalizar(ruta.stem) for ign in ARCHIVOS_IGNORADOS):
                continue
            yield origen, ruta


# ---------------------------------------------------------------------------
# Construcción del archivo de datos
# ---------------------------------------------------------------------------


def clave_registro(r):
    return (r["fecha"], r["dia"], r["carga"], r["esfuerzo"])


def construir(informe):
    ejercicios = {}            # id -> datos en construcción
    ahora = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()

    for origen, ruta in recorrer():
        datos = leer_archivo(ruta)
        if datos is None:
            informe["ignorados"].append(f"{ruta} (no tiene hoja «Ciclo»)")
            continue
        ctx = sede_y_contexto(origen, ruta)
        nombre, nota_nombre = nombre_de_archivo(ruta)
        ej_id = "ej_" + slug(nombre) + ("_" + ctx["sede"].removeprefix("sede_") if ctx["sede"] else "")

        ej = ejercicios.setdefault(ej_id, {
            "id": ej_id, "nombre": nombre, "sedeId": ctx["sede"], "grupo": None,
            "vigente": False, "ciclos": [], "notas": [], "archivos": [],
            "orden": None,
        })
        ej["vigente"] |= ctx["vigente"]
        ej["grupo"] = ej["grupo"] or ctx["grupo"]
        ej["archivos"].append(str(ruta))
        for n in ([nota_nombre] if nota_nombre else []) + datos["notas"]:
            if n not in ej["notas"]:
                ej["notas"].append(n)
        if ctx["vigente"]:
            m = re.match(r"^(\d+)(\.\d+)?\.?\s", ruta.name)
            dia_m = re.match(r"^(\d+)\.", ruta.parent.name)
            if m and dia_m:
                ej.setdefault("enRutina", []).append({
                    "sede": origen["sede"], "dia": int(dia_m.group(1)),
                    "diaNombre": re.sub(r"^\d+\.\s*", "", ruta.parent.name),
                    "orden": float(m.group(1) + (m.group(2) or "")),
                })
        for c in datos["ciclos"]:
            c["vigente"] = ctx["vigente"]
            c["archivo"] = ruta.name
            ej["ciclos"].append(c)

    return [finalizar_ejercicio(ej, informe) for ej in ejercicios.values()], ahora


def finalizar_ejercicio(ej, informe):
    # Tipos de medida: se decide por los bloques que tienen datos.
    con_datos = [c for c in ej["ciclos"] if c["registros"]] or ej["ciclos"]
    tipo_esf = Counter(c["esfuerzo"] for c in con_datos).most_common(1)[0][0]
    tipo_carga = Counter(c["carga"] for c in con_datos).most_common(1)[0][0]
    usa_lectura = any(c["lectura"] for c in ej["ciclos"])
    if ej["grupo"] == "peso-corporal" and tipo_carga == "peso":
        tipo_carga = "pesoCorporal"

    # Quitar ciclos repetidos: la copia antigua contiene versiones anteriores
    # de los mismos archivos, con un subconjunto de los mismos registros.
    ciclos = sorted(ej["ciclos"], key=lambda c: -len(c["registros"]))
    unicos = []
    for c in ciclos:
        claves = {clave_registro(r) for r in c["registros"]}
        if claves and any(claves <= {clave_registro(r) for r in u["registros"]} for u in unicos):
            informe["ciclosDuplicados"] += 1
            continue
        unicos.append(c)

    usados = [c for c in unicos if c["registros"]]
    fechas = lambda c: sorted(r["fecha"] for r in c["registros"] if r["fecha"])
    usados.sort(key=lambda c: (fechas(c)[0] if fechas(c) else "9999", c["archivo"]))

    # El siguiente ciclo ya preparado en el Excel vigente (escalera sin datos).
    preparado = None
    vigentes = [c for c in ej["ciclos"] if c["vigente"]]
    if vigentes:
        ult_archivo = vigentes[-1]["archivo"]
        del_archivo = [c for c in vigentes if c["archivo"] == ult_archivo]
        for i, c in enumerate(del_archivo):
            if not c["registros"] and any(x["registros"] for x in del_archivo[:i]) \
                    and any(v is not None for v in c["escalera"]):
                preparado = c
                break

    ciclos_json = []
    for n, c in enumerate(usados, start=1):
        f = fechas(c)
        ciclos_json.append({
            "n": n,
            "inicio": f[0] if f else None,
            "fin": f[-1] if f else None,
            "escalera": c["escalera"],
            "archivoOrigen": c["archivo"],
            "registros": sorted(c["registros"], key=lambda r: r["dia"]),
        })
    if preparado:
        ciclos_json.append({"n": len(ciclos_json) + 1, "inicio": None, "fin": None,
                            "escalera": preparado["escalera"],
                            "archivoOrigen": preparado["archivo"], "registros": []})

    unidad_carga = {"peso": "kg", "pesoCorporal": "kg", "altura": "cm"}[tipo_carga]
    etiqueta_carga = {"peso": "Peso", "pesoCorporal": "Peso", "altura": "Altura de ladrillo"}[tipo_carga]
    ejercicio = {
        "id": ej["id"],
        "nombre": ej["nombre"],
        "sedeId": ej["sedeId"],
        "grupo": ej["grupo"],
        "archivado": not ej["vigente"],
        "carga": {"tipo": tipo_carga, "unidad": unidad_carga, "etiqueta": etiqueta_carga},
        "esfuerzo": ({"tipo": "tiempo", "unidad": "s", "etiqueta": "Segundos"}
                     if tipo_esf == "tiempo" else
                     {"tipo": "repeticiones", "unidad": "reps", "etiqueta": "Repeticiones"}),
        "lectura": ({"etiqueta": "Kg máq", "formula": "referencia - lectura", "referencia": 95}
                    if usa_lectura else None),
        "formula1RM": "epley",
        "progresion": {
            "tipo": "bilbo",
            "diasPorCiclo": DIAS_POR_CICLO,
            "cicloActual": len(ciclos_json) or None,
            "ciclos": [{k: v for k, v in c.items() if k != "registros"} for c in ciclos_json],
        },
        "notas": "\n".join(ej["notas"]),
        "importado": {"archivos": ej["archivos"]},
    }
    return ejercicio, ciclos_json, ej.get("enRutina", [])


def construir_sesiones(resultado, informe):
    sesiones = {}
    for ejercicio, ciclos, _ in resultado:
        for ciclo in ciclos:
            for r in ciclo["registros"]:
                if not r["fecha"]:
                    informe["sinFecha"].append(f"{ejercicio['nombre']}: ciclo {ciclo['n']}, día {r['dia']}")
                    continue
                sede = ejercicio["sedeId"] or "sede_habitual"
                ses_id = f"ses_{r['fecha'].replace('-', '')}_{sede.removeprefix('sede_')}"
                ses = sesiones.setdefault(ses_id, {
                    "id": ses_id, "fecha": r["fecha"], "sedeId": sede,
                    "rutinaId": None, "diaRutinaId": None, "estado": "terminada",
                    "inicio": None, "fin": None, "ejercicios": [], "notas": "",
                    "origen": "importado-excel",
                })
                if r["fechaAproximada"]:
                    ses["notas"] = "Fecha aproximada"
                # Un mismo día pudo registrar dos días del ciclo: van por separado.
                entrada = next((e for e in ses["ejercicios"]
                                if e["ejercicioId"] == ejercicio["id"] and e["cicloN"] == ciclo["n"]
                                and e["diaCiclo"] == r["dia"]), None)
                if entrada is None:
                    entrada = {"ejercicioId": ejercicio["id"], "cicloN": ciclo["n"],
                               "diaCiclo": r["dia"], "series": [], "notas": ""}
                    ses["ejercicios"].append(entrada)
                entrada["series"].append({
                    "id": f"s{len(entrada['series']) + 1}",
                    "tipo": "bilbo",
                    "carga": r["carga"],
                    "lectura": r["lectura"],
                    "esfuerzo": r["esfuerzo"],
                    "hecha": True,
                })
    return sorted(sesiones.values(), key=lambda s: (s["fecha"], s["sedeId"]))


def construir_rutinas(resultado):
    """Una rutina PLPL por sede, a partir del orden de las carpetas vigentes."""
    por_sede = defaultdict(lambda: defaultdict(list))
    nombres_dia = {}
    for ejercicio, _, en_rutina in resultado:
        for pos in en_rutina:
            por_sede[pos["sede"]][pos["dia"]].append((pos["orden"], ejercicio))
            nombres_dia[(pos["sede"], pos["dia"])] = pos["diaNombre"]

    rutinas = []
    for sede, dias in sorted(por_sede.items()):
        rutina = {"id": f"rut_plpl_2026_{sede.removeprefix('sede_')}",
                  "nombre": "PLPL Bilbo + Heavy Duty", "sedeId": sede,
                  "activa": True, "dias": []}
        for n in sorted(dias):
            ejercicios_dia = []
            for _, ej in sorted(dias[n], key=lambda x: x[0]):
                clave = normalizar(ej["nombre"])
                plan = next((p for p in PLPL_SERIES if p[0] in clave), None)
                if plan:
                    series = [{"tipo": "bilbo", "objetivoEsfuerzo": plan[1]},
                              {"tipo": "intensidad", "tecnica": plan[2], "objetivoEsfuerzo": plan[3]}]
                else:
                    series = [{"tipo": "bilbo", "objetivoEsfuerzo": None}]
                ejercicios_dia.append({"ejercicioId": ej["id"], "opcional": False, "series": series})
            ejercicios_dia.append({"ejercicioId": "ej_cardio", "opcional": True,
                                   "series": [{"tipo": "libre", "objetivoEsfuerzo": [10, 20]}]})
            nombre = nombres_dia[(sede, n)]
            rutina["dias"].append({"id": f"dia_{n}_{slug(nombre)}",
                                   "nombre": f"Día {n} ({nombre})", "ejercicios": ejercicios_dia})
        rutinas.append(rutina)
    return rutinas


CARDIO = {
    "id": "ej_cardio", "nombre": "Cardio", "sedeId": None, "grupo": "cardio",
    "archivado": False,
    "carga": {"tipo": "ninguna"},
    "esfuerzo": {"tipo": "tiempo", "unidad": "min", "etiqueta": "Minutos"},
    "esfuerzoExtra": {"tipo": "distancia", "unidad": "km", "etiqueta": "Distancia", "opcional": True},
    "lectura": None, "formula1RM": None,
    "progresion": {"tipo": "libre"}, "notas": "",
}


def main():
    raiz = Path(__file__).resolve().parent.parent
    salida = raiz / "datos-privados"
    salida.mkdir(exist_ok=True)

    informe = {"ignorados": [], "sinFecha": [], "ciclosDuplicados": 0}
    resultado, ahora = construir(informe)
    sesiones = construir_sesiones(resultado, informe)
    rutinas = construir_rutinas(resultado)
    ejercicios = sorted((e for e, _, _ in resultado),
                        key=lambda e: (e["archivado"], e["sedeId"] or "", e["grupo"] or "~", e["nombre"]))

    datos = {
        "version": VERSION_ESQUEMA,
        "creado": ahora,
        "actualizado": ahora,
        "revision": 1,
        "origen": "importacion-excel",
        "perfil": {"nombre": "Dan", "correo": None, "pesoCorporalKg": None,
                   "unidadPeso": "kg", "sedePorDefecto": "sede_habitual", "tema": "sistema"},
        "sedes": SEDES,
        "ejercicios": ejercicios + [CARDIO],
        "rutinas": rutinas,
        "sesiones": sesiones,
    }
    (salida / "entrenamiento-importado.json").write_text(
        json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")

    escribir_informe(salida / "informe-importacion.txt", datos, informe)
    print((salida / "informe-importacion.txt").read_text(encoding="utf-8"))


def escribir_informe(ruta, datos, informe):
    ej = datos["ejercicios"]
    series = sum(len(e["series"]) for s in datos["sesiones"] for e in s["ejercicios"])
    fechas = [s["fecha"] for s in datos["sesiones"]]
    lineas = [
        "INFORME DE IMPORTACIÓN",
        "=" * 60,
        f"Ejercicios: {len(ej)} ({sum(not e['archivado'] for e in ej)} vigentes, "
        f"{sum(e['archivado'] for e in ej)} archivados)",
        f"Sesiones (días entrenados): {len(datos['sesiones'])}",
        f"Series registradas: {series}",
        f"Periodo: {min(fechas)} a {max(fechas)}" if fechas else "Periodo: sin fechas",
        f"Ciclos duplicados descartados (copia antigua): {informe['ciclosDuplicados']}",
        f"Registros sin fecha (no se convierten en sesión): {len(informe['sinFecha'])}",
        "",
    ]
    nombre_sede = {s["id"]: s["nombre"] for s in datos["sedes"]}
    for archivado in (False, True):
        lineas.append("EJERCICIOS " + ("ARCHIVADOS" if archivado else "VIGENTES"))
        lineas.append("-" * 60)
        for e in (x for x in ej if x["archivado"] == archivado):
            prog = e["progresion"]
            ciclos = prog.get("ciclos", [])
            n_reg = sum(len(s2["series"]) for s in datos["sesiones"] for s2 in s["ejercicios"]
                        if s2["ejercicioId"] == e["id"])
            medida = e["esfuerzo"]["etiqueta"]
            if e.get("lectura"):
                medida += " + Kg máq"
            if e["carga"]["tipo"] == "altura":
                medida += " (altura)"
            lineas.append(f"  {e['nombre']:<42} {nombre_sede.get(e['sedeId'], 'Cualquiera'):<10} "
                          f"{len(ciclos):>2} ciclos {n_reg:>4} series  [{medida}]")
        lineas.append("")
    for s in datos["rutinas"]:
        lineas.append(f"RUTINA {s['nombre']} ({nombre_sede[s['sedeId']]})")
        lineas.append("-" * 60)
        nombres = {e["id"]: e["nombre"] for e in ej}
        for d in s["dias"]:
            lineas.append(f"  {d['nombre']}: " + ", ".join(nombres[x["ejercicioId"]] for x in d["ejercicios"]))
        lineas.append("")
    if informe["sinFecha"]:
        lineas.append("REGISTROS SIN FECHA")
        lineas.append("-" * 60)
        lineas += [f"  {x}" for x in informe["sinFecha"]]
        lineas.append("")
    if informe["ignorados"]:
        lineas.append("ARCHIVOS IGNORADOS")
        lineas += [f"  {x}" for x in informe["ignorados"]]
    ruta.write_text("\n".join(lineas), encoding="utf-8")


if __name__ == "__main__":
    sys.exit(main())
